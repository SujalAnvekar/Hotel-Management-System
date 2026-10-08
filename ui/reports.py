import tkinter as tk
from tkinter import ttk, messagebox
import os

from openpyxl import Workbook
from openpyxl.styles import Font

from database.connection import get_connection


class ReportsPage:

    def __init__(self, parent):

        self.parent = parent

        # Create reports folder automatically
        self.reports_folder = os.path.join(
            os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
            "reports"
        )

        if not os.path.exists(self.reports_folder):
            os.makedirs(self.reports_folder)

        self.create_widgets()
        self.load_sales_report()

    # CREATE WIDGETS
    
    def create_widgets(self):

        title = tk.Label(
            self.parent,
            text="Reports",
            font=("Arial", 20, "bold")
        )

        title.pack(
            anchor="w",
            padx=20,
            pady=20
        )

        # REPORT BUTTONS

        button_frame = tk.Frame(
            self.parent
        )

        button_frame.pack(
            fill="x",
            padx=20,
            pady=10
        )

        # Sales Report
        tk.Button(
            button_frame,
            text="Sales Report",
            width=15,
            command=self.load_sales_report
        ).pack(
            side="left",
            padx=5
        )

        # Orders Report
        tk.Button(
            button_frame,
            text="Orders Report",
            width=15,
            command=self.load_orders_report
        ).pack(
            side="left",
            padx=5
        )

        # Payment Report
        tk.Button(
            button_frame,
            text="Payment Report",
            width=15,
            command=self.load_payment_report
        ).pack(
            side="left",
            padx=5
        )

        # Inventory Report
        tk.Button(
            button_frame,
            text="Inventory Report",
            width=15,
            command=self.load_inventory_report
        ).pack(
            side="left",
            padx=5
        )

        # SUMMARY
        summary_frame = tk.Frame(
            self.parent
        )

        summary_frame.pack(
            fill="x",
            padx=20,
            pady=15
        )

        self.total_bills_label = tk.Label(
            summary_frame,
            text="Total Bills: 0",
            font=("Arial", 12, "bold")
        )

        self.total_bills_label.pack(
            side="left",
            padx=20
        )

        self.total_sales_label = tk.Label(
            summary_frame,
            text="Total Sales: ₹0.00",
            font=("Arial", 12, "bold")
        )

        self.total_sales_label.pack(
            side="left",
            padx=20
        )

        # REPORT TABLE

        table_frame = tk.Frame(
            self.parent
        )

        table_frame.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=10
        )

        columns = (
            "BillID",
            "OrderID",
            "BillDate",
            "SubTotal",
            "TaxAmount",
            "DiscountAmount",
            "TotalAmount"
        )

        self.tree = ttk.Treeview(
            table_frame,
            columns=columns,
            show="headings"
        )

        self.tree.heading(
            "BillID",
            text="Bill ID"
        )

        self.tree.heading(
            "OrderID",
            text="Order ID"
        )

        self.tree.heading(
            "BillDate",
            text="Bill Date"
        )

        self.tree.heading(
            "SubTotal",
            text="Subtotal"
        )

        self.tree.heading(
            "TaxAmount",
            text="Tax"
        )

        self.tree.heading(
            "DiscountAmount",
            text="Discount"
        )

        self.tree.heading(
            "TotalAmount",
            text="Total"
        )

        self.tree.column(
            "BillID",
            width=70
        )

        self.tree.column(
            "OrderID",
            width=80
        )

        self.tree.column(
            "BillDate",
            width=150
        )

        self.tree.column(
            "SubTotal",
            width=100
        )

        self.tree.column(
            "TaxAmount",
            width=100
        )

        self.tree.column(
            "DiscountAmount",
            width=100
        )

        self.tree.column(
            "TotalAmount",
            width=120
        )

        self.tree.pack(
            fill="both",
            expand=True
        )


    # SALES REPORT


    def load_sales_report(self):

        for item in self.tree.get_children():
            self.tree.delete(item)

        columns = (
            "BillID",
            "OrderID",
            "BillDate",
            "SubTotal",
            "TaxAmount",
            "DiscountAmount",
            "TotalAmount"
        )

        self.tree["columns"] = columns

        self.tree.heading(
            "BillID",
            text="Bill ID"
        )

        self.tree.heading(
            "OrderID",
            text="Order ID"
        )

        self.tree.heading(
            "BillDate",
            text="Bill Date"
        )

        self.tree.heading(
            "SubTotal",
            text="Subtotal"
        )

        self.tree.heading(
            "TaxAmount",
            text="Tax"
        )

        self.tree.heading(
            "DiscountAmount",
            text="Discount"
        )

        self.tree.heading(
            "TotalAmount",
            text="Total"
        )

        self.tree.column(
            "BillID",
            width=70
        )

        self.tree.column(
            "OrderID",
            width=80
        )

        self.tree.column(
            "BillDate",
            width=150
        )

        self.tree.column(
            "SubTotal",
            width=100
        )

        self.tree.column(
            "TaxAmount",
            width=100
        )

        self.tree.column(
            "DiscountAmount",
            width=100
        )

        self.tree.column(
            "TotalAmount",
            width=120
        )

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT
                    BillID,
                    OrderID,
                    BillDate,
                    SubTotal,
                    TaxAmount,
                    DiscountAmount,
                    TotalAmount
                FROM Bills
                ORDER BY BillID DESC
            """)

            rows = cursor.fetchall()

            connection.close()

            total_bills = 0
            total_sales = 0

            for row in rows:

                subtotal = float(row[3])
                tax_amount = float(row[4])
                discount_amount = float(row[5])
                total_amount = float(row[6])

                self.tree.insert(
                    "",
                    "end",
                    values=(
                        row[0],
                        row[1],
                        row[2],
                        f"₹{subtotal:.2f}",
                        f"₹{tax_amount:.2f}",
                        f"₹{discount_amount:.2f}",
                        f"₹{total_amount:.2f}"
                    )
                )

                total_bills += 1
                total_sales += total_amount

            self.total_bills_label.config(
                text=f"Total Bills: {total_bills}"
            )

            self.total_sales_label.config(
                text=f"Total Sales: ₹{total_sales:.2f}"
            )

            # Automatically save Excel
            self.export_sales_to_excel(rows)

        except Exception as e:

            messagebox.showerror(
                "Error",
                str(e)
            )

    # EXPORT SALES TO EXCEL

    def export_sales_to_excel(self, rows):

        try:

            file_path = os.path.join(
                self.reports_folder,
                "Sales_Report.xlsx"
            )

            workbook = Workbook()

            worksheet = workbook.active
            worksheet.title = "Sales Report"

            headers = [
                "Bill ID",
                "Order ID",
                "Bill Date",
                "Subtotal",
                "Tax Amount",
                "Discount Amount",
                "Total Amount"
            ]

            # Header row
            for column, header in enumerate(headers, start=1):

                cell = worksheet.cell(
                    row=1,
                    column=column,
                    value=header
                )

                cell.font = Font(
                    bold=True
                )

            # Data
            for row_number, row in enumerate(rows, start=2):

                worksheet.cell(
                    row=row_number,
                    column=1,
                    value=row[0]
                )

                worksheet.cell(
                    row=row_number,
                    column=2,
                    value=row[1]
                )

                worksheet.cell(
                    row=row_number,
                    column=3,
                    value=row[2]
                )

                worksheet.cell(
                    row=row_number,
                    column=4,
                    value=float(row[3])
                )

                worksheet.cell(
                    row=row_number,
                    column=5,
                    value=float(row[4])
                )

                worksheet.cell(
                    row=row_number,
                    column=6,
                    value=float(row[5])
                )

                worksheet.cell(
                    row=row_number,
                    column=7,
                    value=float(row[6])
                )

            # Column widths
            worksheet.column_dimensions["A"].width = 12
            worksheet.column_dimensions["B"].width = 12
            worksheet.column_dimensions["C"].width = 25
            worksheet.column_dimensions["D"].width = 15
            worksheet.column_dimensions["E"].width = 15
            worksheet.column_dimensions["F"].width = 18
            worksheet.column_dimensions["G"].width = 15

            # Currency formatting
            for row_number in range(2, len(rows) + 2):

                worksheet.cell(
                    row=row_number,
                    column=4
                ).number_format = '₹#,##0.00'

                worksheet.cell(
                    row=row_number,
                    column=5
                ).number_format = '₹#,##0.00'

                worksheet.cell(
                    row=row_number,
                    column=6
                ).number_format = '₹#,##0.00'

                worksheet.cell(
                    row=row_number,
                    column=7
                ).number_format = '₹#,##0.00'

            workbook.save(file_path)

        except Exception as e:

            print("Sales Excel Error:", e)

    # ORDERS REPORT

    def load_orders_report(self):

        for item in self.tree.get_children():
            self.tree.delete(item)

        columns = (
            "OrderID",
            "Customer",
            "OrderDate",
            "Status",
            "TotalAmount"
        )

        self.tree["columns"] = columns

        self.tree.heading(
            "OrderID",
            text="Order ID"
        )

        self.tree.heading(
            "Customer",
            text="Customer"
        )

        self.tree.heading(
            "OrderDate",
            text="Order Date"
        )

        self.tree.heading(
            "Status",
            text="Status"
        )

        self.tree.heading(
            "TotalAmount",
            text="Total Amount"
        )

        self.tree.column(
            "OrderID",
            width=80
        )

        self.tree.column(
            "Customer",
            width=180
        )

        self.tree.column(
            "OrderDate",
            width=160
        )

        self.tree.column(
            "Status",
            width=120
        )

        self.tree.column(
            "TotalAmount",
            width=120
        )

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT
                    o.OrderID,
                    c.full_name,
                    o.OrderDate,
                    o.Status,
                    o.TotalAmount
                FROM Orders o
                INNER JOIN Customers c
                    ON o.customer_id = c.customer_id
                ORDER BY o.OrderID DESC
            """)

            rows = cursor.fetchall()

            connection.close()

            total_orders = 0
            total_amount = 0

            for row in rows:

                order_total = float(row[4])

                self.tree.insert(
                    "",
                    "end",
                    values=(
                        row[0],
                        row[1],
                        row[2],
                        row[3],
                        f"₹{order_total:.2f}"
                    )
                )

                total_orders += 1
                total_amount += order_total

            self.total_bills_label.config(
                text=f"Total Orders: {total_orders}"
            )

            self.total_sales_label.config(
                text=f"Order Amount: ₹{total_amount:.2f}"
            )

            # Automatically save Excel
            self.export_orders_to_excel(rows)

        except Exception as e:

            messagebox.showerror(
                "Error",
                str(e)
            )

    # EXPORT ORDERS TO EXCEL

    def export_orders_to_excel(self, rows):

        try:

            file_path = os.path.join(
                self.reports_folder,
                "Orders_Report.xlsx"
            )

            workbook = Workbook()

            worksheet = workbook.active
            worksheet.title = "Orders Report"

            headers = [
                "Order ID",
                "Customer",
                "Order Date",
                "Status",
                "Total Amount"
            ]

            # Header row
            for column, header in enumerate(headers, start=1):

                cell = worksheet.cell(
                    row=1,
                    column=column,
                    value=header
                )

                cell.font = Font(
                    bold=True
                )

            # Data
            for row_number, row in enumerate(rows, start=2):

                worksheet.cell(
                    row=row_number,
                    column=1,
                    value=row[0]
                )

                worksheet.cell(
                    row=row_number,
                    column=2,
                    value=row[1]
                )

                worksheet.cell(
                    row=row_number,
                    column=3,
                    value=row[2]
                )

                worksheet.cell(
                    row=row_number,
                    column=4,
                    value=row[3]
                )

                worksheet.cell(
                    row=row_number,
                    column=5,
                    value=float(row[4])
                )

            # Column widths
            worksheet.column_dimensions["A"].width = 12
            worksheet.column_dimensions["B"].width = 25
            worksheet.column_dimensions["C"].width = 25
            worksheet.column_dimensions["D"].width = 15
            worksheet.column_dimensions["E"].width = 18

            # Currency formatting
            for row_number in range(2, len(rows) + 2):

                worksheet.cell(
                    row=row_number,
                    column=5
                ).number_format = '₹#,##0.00'

            workbook.save(file_path)

        except Exception as e:

            print("Orders Excel Error:", e)

    # PAYMENT REPORT

    def load_payment_report(self):

        for item in self.tree.get_children():
            self.tree.delete(item)

        columns = (
            "PaymentID",
            "BillID",
            "OrderID",
            "PaymentDate",
            "PaymentMethod",
            "Amount"
        )

        self.tree["columns"] = columns

        self.tree.heading(
            "PaymentID",
            text="Payment ID"
        )

        self.tree.heading(
            "BillID",
            text="Bill ID"
        )

        self.tree.heading(
            "OrderID",
            text="Order ID"
        )

        self.tree.heading(
            "PaymentDate",
            text="Payment Date"
        )

        self.tree.heading(
            "PaymentMethod",
            text="Payment Method"
        )

        self.tree.heading(
            "Amount",
            text="Amount"
        )

        self.tree.column(
            "PaymentID",
            width=90
        )

        self.tree.column(
            "BillID",
            width=80
        )

        self.tree.column(
            "OrderID",
            width=80
        )

        self.tree.column(
            "PaymentDate",
            width=160
        )

        self.tree.column(
            "PaymentMethod",
            width=130
        )

        self.tree.column(
            "Amount",
            width=120
        )

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT
                    p.PaymentID,
                    p.BillID,
                    b.OrderID,
                    p.PaymentDate,
                    p.PaymentMethod,
                    p.Amount
                FROM Payments p
                INNER JOIN Bills b
                    ON p.BillID = b.BillID
                ORDER BY p.PaymentID DESC
            """)

            rows = cursor.fetchall()

            connection.close()

            total_payments = 0
            total_amount = 0

            for row in rows:

                amount = float(row[5])

                self.tree.insert(
                    "",
                    "end",
                    values=(
                        row[0],
                        row[1],
                        row[2],
                        row[3],
                        row[4],
                        f"₹{amount:.2f}"
                    )
                )

                total_payments += 1
                total_amount += amount

            self.total_bills_label.config(
                text=f"Total Payments: {total_payments}"
            )

            self.total_sales_label.config(
                text=f"Payment Amount: ₹{total_amount:.2f}"
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                str(e)
            )


    # INVENTORY REPORT

    def load_inventory_report(self):

        for item in self.tree.get_children():
            self.tree.delete(item)

        columns = (
            "InventoryID",
            "ItemName",
            "Unit",
            "Quantity",
            "MinimumQuantity",
            "Status"
        )

        self.tree["columns"] = columns

        self.tree.heading(
            "InventoryID",
            text="Inventory ID"
        )

        self.tree.heading(
            "ItemName",
            text="Item Name"
        )

        self.tree.heading(
            "Unit",
            text="Unit"
        )

        self.tree.heading(
            "Quantity",
            text="Quantity"
        )

        self.tree.heading(
            "MinimumQuantity",
            text="Minimum Quantity"
        )

        self.tree.heading(
            "Status",
            text="Status"
        )

        self.tree.column(
            "InventoryID",
            width=100
        )

        self.tree.column(
            "ItemName",
            width=180
        )

        self.tree.column(
            "Unit",
            width=100
        )

        self.tree.column(
            "Quantity",
            width=120
        )

        self.tree.column(
            "MinimumQuantity",
            width=140
        )

        self.tree.column(
            "Status",
            width=120
        )

        try:

            connection = get_connection()
            cursor = connection.cursor()

            cursor.execute("""
                SELECT
                    InventoryID,
                    ItemName,
                    Unit,
                    Quantity,
                    MinimumQuantity
                FROM Inventory
                ORDER BY InventoryID DESC
            """)

            rows = cursor.fetchall()

            connection.close()

            total_items = 0
            low_stock_items = 0

            for row in rows:

                inventory_id = row[0]
                item_name = row[1]
                unit = row[2]
                quantity = float(row[3])
                minimum_quantity = float(row[4])

                if quantity <= minimum_quantity:
                    status = "Low Stock"
                    low_stock_items += 1
                else:
                    status = "Available"

                self.tree.insert(
                    "",
                    "end",
                    values=(
                        inventory_id,
                        item_name,
                        unit,
                        f"{quantity:.3f}",
                        f"{minimum_quantity:.3f}",
                        status
                    )
                )

                total_items += 1

            self.total_bills_label.config(
                text=f"Total Items: {total_items}"
            )

            self.total_sales_label.config(
                text=f"Low Stock: {low_stock_items}"
            )

        except Exception as e:

            messagebox.showerror(
                "Error",
                str(e)
            )